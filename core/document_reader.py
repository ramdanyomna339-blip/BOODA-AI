"""Bounded, read-only extraction from explicitly selected project documents."""
import csv
import hashlib
import io
from itertools import islice
from pathlib import Path
import zipfile

from core.laptop import BoodaLaptopConnection

TYPES = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md", ".png", ".jpg", ".jpeg"}


def selected_file(root, relative):
    connection = BoodaLaptopConnection(root)
    # Reuse the scoped path guards while checking file types here.
    path = connection._path(relative, directory=True)
    if not path.is_file() or path.suffix.lower() not in TYPES or path.stat().st_size > 20_000_000:
        raise ValueError("اختر ملفًا مدعومًا لا يتجاوز 20 ميجابايت داخل المجلد المسموح.")
    return path


def check_zip(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        rows = archive.infolist()
        if (len(rows) > 5000 or sum(row.file_size for row in rows) > 50_000_000
                or any(row.file_size > 10_000_000 or row.flag_bits & 1 for row in rows)):
            raise ValueError("الملف المضغوط كبير أو مشفر؛ استخدم نسخة أصغر.")


def read_document(root, relative):
    path = selected_file(root, relative)
    with path.open("rb") as stream:
        raw = stream.read(20_000_001)
    if len(raw) > 20_000_000:
        raise ValueError("الملف تجاوز الحد أثناء القراءة.")
    suffix = path.suffix.lower()
    parts, truncated = [], False
    if suffix in {".docx", ".xlsx"}:
        check_zip(raw)
    if suffix == ".docx":
        from defusedxml.ElementTree import fromstring
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            document = fromstring(archive.read("word/document.xml"))
            for paragraph in document.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                parts.append("".join(node.text or "" for node in paragraph.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets[:10]:
                parts.append("[" + sheet.title + "]")
                for row in sheet.iter_rows(max_row=200, max_col=30, values_only=True):
                    parts.append("\t".join(str(value)[:1000] if value is not None else "" for value in row))
                truncated |= sheet.max_row is None or sheet.max_row > 200 or (sheet.max_column or 0) > 30
            truncated |= len(workbook.worksheets) > 10
        finally:
            workbook.close()
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("افتح نسخة غير مشفرة بإذنك؛ لا يتم تخمين كلمات المرور.")
        for number, page in enumerate(islice(reader.pages, 40), 1):
            parts.append(f"[{number}]\n" + (page.extract_text() or ""))
        truncated = len(reader.pages) > 40
    elif suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image
        from core.image_text import read_image_text
        with Image.open(io.BytesIO(raw)) as image:
            parts.append(read_image_text(image)["text"])
    elif suffix == ".csv":
        rows = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
        for row in islice(rows, 201):
            if len(parts) == 200:
                truncated = True
                break
            truncated |= len(row) > 30
            parts.append("\t".join(value[:1000] for value in row[:30]))
    else:
        parts.append(raw.decode("utf-8-sig"))
    text = "\n".join(parts)
    return {"source": relative, "text": text[:30000], "truncated": truncated or len(text) > 30000,
            "sha256": hashlib.sha256(raw).hexdigest(), "untrusted": True,
            "requires_review": True, "formulas_evaluated": False,
            "note": "PDF المصور يحتاج OCR منفصلًا؛ الصيغ تُعرض كنص ولا تُحسب."}
